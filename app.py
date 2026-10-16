
import io
import json
import re

import pandas as pd
import streamlit as st
from groq import Groq
from pypdf import PdfReader
from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer
)
from reportlab.lib.styles import getSampleStyleSheet


st.set_page_config(
    page_title="QuizCraft AI",
    page_icon="📝",
    layout="wide"
)

st.title("📝 QuizCraft AI")
st.caption("Turn your study material into editable quizzes in minutes.")

st.info(
    "Prototype version: generated questions should be reviewed "
    "by a teacher before being used in an exam."
)


def extract_text(uploaded_file):
    """Extract text from PDF, DOCX or TXT."""
    name = uploaded_file.name.lower()

    if name.endswith(".pdf"):
        reader = PdfReader(uploaded_file)
        pages = []
        for page in reader.pages:
            pages.append(page.extract_text() or "")
        return "\n".join(pages)

    if name.endswith(".docx"):
        document = Document(uploaded_file)
        return "\n".join(
            paragraph.text for paragraph in document.paragraphs
        )

    if name.endswith(".txt"):
        return uploaded_file.getvalue().decode(
            "utf-8", errors="replace"
        )

    raise ValueError("Unsupported file format.")


def generate_quiz(
    source_text,
    subject,
    level,
    difficulty,
    question_type,
    number_of_questions,
    language,
):
    client = Groq(api_key=st.secrets["GROQ_API_KEY"])

    prompt = f"""
Create a high-quality quiz using ONLY the learning material below.
Do not invent facts that are not supported by the material.
If the material is insufficient, make fewer questions rather than
inventing unsupported answers.

Subject: {subject}
Learner level: {level}
Difficulty: {difficulty}
Question format: {question_type}
Requested number of questions: {number_of_questions}
Language: {language}

Return valid JSON only, using this exact structure:
{{
  "questions": [
    {{
      "question": "Question text",
      "options": ["Option A", "Option B", "Option C", "Option D"],
      "correct_answer": "Exact correct option or answer",
      "explanation": "Brief explanation grounded in the material"
    }}
  ]
}}

Rules:
- For multiple choice, provide exactly four options and one correct answer.
- For true/false, provide options ["True", "False"].
- For short answer, use an empty options list.
- The correct_answer must match an option exactly for MCQs and true/false.
- Avoid duplicate questions.
- Keep explanations concise.
- Do not include markdown fences around the JSON.

LEARNING MATERIAL:
{source_text}
"""

    response = client.chat.completions.create(
        model="openai/gpt-oss-20b",
        messages=[
            {
                "role": "system",
                "content": (
                    "You are an educational assessment assistant. "
                    "Follow the supplied source material and output "
                    "valid JSON only."
                ),
            },
            {"role": "user", "content": prompt},
        ],
        temperature=0.2,
        max_completion_tokens=6000,
        response_format={"type": "json_object"},
    )

    content = response.choices[0].message.content
    data = json.loads(content)
    questions = data.get("questions", [])

    if not isinstance(questions, list) or not questions:
        raise ValueError("The AI returned no usable questions.")

    clean_questions = []
    for item in questions:
        if not isinstance(item, dict):
            continue

        question = str(item.get("question", "")).strip()
        if not question:
            continue

        options = item.get("options", [])
        if not isinstance(options, list):
            options = []

        clean_questions.append({
            "Question": question,
            "Option A": str(options[0]) if len(options) > 0 else "",
            "Option B": str(options[1]) if len(options) > 1 else "",
            "Option C": str(options[2]) if len(options) > 2 else "",
            "Option D": str(options[3]) if len(options) > 3 else "",
            "Correct answer": str(
                item.get("correct_answer", "")
            ).strip(),
            "Explanation": str(
                item.get("explanation", "")
            ).strip(),
        })

    if not clean_questions:
        raise ValueError("No valid questions were returned.")

    return pd.DataFrame(clean_questions)


def make_docx(df, title):
    document = Document()
    document.add_heading(title, 0)

    for index, row in df.iterrows():
        document.add_heading(
            f"Question {index + 1}", level=2
        )
        document.add_paragraph(str(row["Question"]))

        for column in ["Option A", "Option B", "Option C", "Option D"]:
            value = str(row[column]).strip()
            if value and value.lower() != "nan":
                document.add_paragraph(
                    f"{column}: {value}"
                )

        document.add_paragraph(
            f"Answer: {row['Correct answer']}"
        )
        explanation = str(row["Explanation"]).strip()
        if explanation and explanation.lower() != "nan":
            document.add_paragraph(f"Explanation: {explanation}")

    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def make_pdf(df, title):
    output = io.BytesIO()
    pdf = SimpleDocTemplate(
        output,
        pagesize=letter,
        rightMargin=45,
        leftMargin=45,
        topMargin=45,
        bottomMargin=45,
    )
    styles = getSampleStyleSheet()
    story = [Paragraph(title, styles["Title"]), Spacer(1, 16)]

    for index, row in df.iterrows():
        story.append(Paragraph(
            f"<b>Question {index + 1}</b>",
            styles["Heading3"],
        ))
        story.append(Paragraph(
            str(row["Question"]), styles["BodyText"]
        ))

        for column in ["Option A", "Option B", "Option C", "Option D"]:
            value = str(row[column]).strip()
            if value and value.lower() != "nan":
                story.append(Paragraph(
                    f"{column}: {value}", styles["BodyText"]
                ))

        story.append(Paragraph(
            f"<b>Answer:</b> {row['Correct answer']}",
            styles["BodyText"],
        ))

        explanation = str(row["Explanation"]).strip()
        if explanation and explanation.lower() != "nan":
            story.append(Paragraph(
                f"<b>Explanation:</b> {explanation}",
                styles["BodyText"],
            ))

        story.append(Spacer(1, 12))

    pdf.build(story)
    return output.getvalue()


with st.sidebar:
    st.header("Quiz settings")
    subject = st.text_input("Subject", value="General Science")
    level = st.selectbox(
        "Learner level",
        [
            "Primary school",
            "Middle school",
            "Secondary school",
            "College",
            "University",
            "Professional training",
        ],
    )
    difficulty = st.selectbox(
        "Difficulty", ["Easy", "Medium", "Hard", "Mixed"]
    )
    question_type = st.selectbox(
        "Question format",
        ["Multiple choice", "True/False", "Short answer"],
    )
    number_of_questions = st.slider(
        "Number of questions", min_value=3, max_value=15, value=5
    )
    language = st.selectbox(
        "Quiz language", ["English", "Urdu"]
    )

uploaded_file = st.file_uploader(
    "Upload study material",
    type=["pdf", "docx", "txt"],
    help="Use a text-based PDF. Scanned images may need OCR.",
)

pasted_text = st.text_area(
    "Or paste your notes here",
    height=180,
    placeholder="Paste your lesson, notes, chapter excerpt or study material...",
)

source_text = ""
source_name = "Pasted study material"

if uploaded_file is not None:
    try:
        source_text = extract_text(uploaded_file)
        source_name = uploaded_file.name
    except Exception as exc:
        st.error(f"Could not read the uploaded file: {exc}")

if pasted_text.strip():
    source_text = pasted_text.strip()
    source_name = "Pasted study material"

if source_text:
    st.caption(
        f"Source: {source_name} · "
        f"{len(source_text):,} characters extracted"
    )

    if len(source_text) > 30000:
        st.warning(
            "Only the first 30,000 characters will be used in this prototype."
        )
        source_text = source_text[:30000]

if "quiz_df" not in st.session_state:
    st.session_state.quiz_df = None

if st.button(
    "✨ Generate quiz",
    type="primary",
    use_container_width=True,
):
    if not source_text.strip():
        st.warning("Upload a file or paste study material first.")
    elif len(source_text.strip()) < 100:
        st.warning(
            "Please provide at least 100 characters of learning material."
        )
    elif not st.secrets.get("GROQ_API_KEY"):
        st.error("GROQ_API_KEY is missing from the app secrets.")
    else:
        try:
            with st.spinner("Reading material and creating questions..."):
                new_df = generate_quiz(
                    source_text=source_text,
                    subject=subject,
                    level=level,
                    difficulty=difficulty,
                    question_type=question_type,
                    number_of_questions=number_of_questions,
                    language=language,
                )
            st.session_state.quiz_df = new_df
            st.success(f"Generated {len(new_df)} questions.")
        except Exception as exc:
            st.error(
                "Quiz generation failed. Check your Groq API key, "
                "model access and API limits. Details: "
                f"{type(exc).__name__}: {exc}"
            )

if st.session_state.quiz_df is not None:
    st.divider()
    st.header("Review and edit your quiz")
    st.caption(
        "Click any cell to edit it. Check every answer before distributing the quiz."
    )

    edited_df = st.data_editor(
        st.session_state.quiz_df,
        num_rows="dynamic",
        hide_index=True,
        use_container_width=True,
        key="quiz_editor",
        column_config={
            "Question": st.column_config.TextColumn(
                "Question", width="large", required=True
            ),
            "Explanation": st.column_config.TextColumn(
                "Explanation", width="large"
            ),
        },
    )
    st.session_state.quiz_df = edited_df

    quiz_title = st.text_input(
        "Quiz title", value=f"{subject} Quiz"
    )

    col1, col2 = st.columns(2)

    with col1:
        st.download_button(
            "Download Word (.docx)",
            data=make_docx(edited_df, quiz_title),
            file_name="quiz.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            use_container_width=True,
        )

    with col2:
        st.download_button(
            "Download PDF",
            data=make_pdf(edited_df, quiz_title),
            file_name="quiz.pdf",
            mime="application/pdf",
            use_container_width=True,
        )

    if st.button("Clear quiz"):
        st.session_state.quiz_df = None
        st.rerun()

st.divider()
st.caption(
    "QuizCraft AI prototype · Review generated content before use."
)
